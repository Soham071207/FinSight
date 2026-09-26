import docx
import re
import os

md_path = r"C:\Users\soham\.gemini\antigravity-ide\brain\0b55af0f-03cf-4f58-b4e9-9034d092ce9b\patent_draft.md"
docx_path = r"c:\Users\soham\Desktop\final1 asep2\FinSight_Patent_Draft.docx"

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    try:
        with open(md_path, 'r', encoding='utf-8') as f:
            content = f.read()
            
        lines = content.split('\n')
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            if line.startswith('# '):
                doc.add_heading(line[2:], 1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], 2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], 3)
            elif line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                # basic bold handling
                line_content = line[2:]
                parts = re.split(r'(\*\*.*?\*\*)', line_content)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)
            else:
                p = doc.add_paragraph()
                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    else:
                        p.add_run(part)
        
        doc.save(docx_path)
        print(f"Successfully saved to {docx_path}")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    convert_md_to_docx(md_path, docx_path)
