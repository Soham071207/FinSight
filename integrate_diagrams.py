import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = r"C:\Users\soham\Desktop\final1 asep2"
DOC_PATH = os.path.join(ROOT, "FinSight_SCI_Journal_Paper.docx")
PLOTS_DIR = os.path.join(ROOT, "STOCK_experimental", "plots")
OUT_DOC = os.path.join(ROOT, "FinSight_SCI_Journal_Paper_Final.docx")

# Mappings: Text to find -> list of tuples (img_name, caption, explanation)
MAPPINGS = {
    "B. Temporal Convolutional Network with Multi-Head Self-Attention (TCN-MHA)": [
        ("fig_system_arch.png", 
         "Fig. 1. System Architecture Schematic Flowchart outlining the multi-modal fusion of structural, temporal, and spatial features.", 
         "As depicted in Fig. 1, the proposed architecture introduces a multi-modal synchronization approach. By combining localized temporal features extracted by the TCN-MHA with broader spatial context from the ST-GCN, the system seeks to anchor individual asset predictions to cross-asset sector dependencies.")
    ],
    "C. Hierarchical Confidence-Gated Sentiment Fusion": [
        ("fig_sentiment_overlay.png", 
         "Fig. 2. FinBERT sentiment momentum overlaid against asset price.",
         "Fig. 2 illustrates the Hierarchical Confidence-Gated Sentiment engine. By mapping FinBERT's raw logit confidence against historical price action, the model dynamically filters sentiment influence, potentially reducing the impact of NLP hallucinations prior to regime shifts.")
    ],
    "E. Spatial-Temporal Graph Convolutional Network (ST-GCN)": [
        ("fig_gcn_heatmap.png",
         "Fig. 3. Dynamic adjacency matrix learned by the Spatial-Temporal GCN.",
         "The spatial heat map in Fig. 3 visualizes the application of Graph Neural Networks to equity pairs trading. The Bilinear Graph Attention Network (GAT) dynamically learns structural adjacency weights, allowing the model to incorporate sector-wide momentum shocks into individual asset predictions without manual sector classification.")
    ],
    "F. Bayesian Meta-Learner with Dual-Perturbation MC Simulation": [
        ("fig_bayesian_violin.png",
         "Fig. 4. Epistemic uncertainty distribution via Monte Carlo dropout for correct vs. incorrect classifications.",
         "The violin plot in Fig. 4 illustrates the isolation of epistemic uncertainty. By running 30 forward Monte Carlo passes, the Bayesian Meta-Learner distinguishes model doubt (epistemic) from inherent market noise (aleatoric). The distribution suggests that incorrect predictions exhibit higher variance, supporting the use of an uncertainty threshold as a structural veto mechanism.")
    ],
    "G. TS-TCC Self-Supervised Pre-Training": [
        ("fig_tstcc_tsne.png",
         "Fig. 5. t-SNE projection demonstrating latent state separation post TS-TCC pre-training.",
         "Fig. 5 provides a qualitative view of the Temporal Contrastive pre-training phase. By applying TS-TCC to unlabeled historical price action, the encoder learns robust representations. The t-SNE projection indicates a degree of topological separation between latent Bull and Bear states, consistent with self-supervised clustering objectives.")
    ],
    "H. Walk-Forward Online Learning for Concept Drift Adaptation": [
        ("fig_walk_forward.png",
         "Fig. 6. Walk-Forward Online Learning Methodology addressing financial concept drift.",
         "Fig. 6 outlines the concept drift adaptation methodology. By utilizing a rolling 120-day validation window and dynamically appending new regression trees to the LightGBM decision layer via the init_model protocol, the architecture incrementally adapts to changing market conditions.")
    ],
    "A. Overall Performance Comparison": [
        ("fig_bar_trend.png",
         "Fig. 7. Comparative bar chart with trend lines summarizing Accuracy and AUC-ROC.",
         "The comparative performance chart in Fig. 7 summarizes the empirical results of the FinSight ensemble. By showing modest improvements over isolated baseline models like XGBoost and Random Forest, the results are consistent with the hypothesis that fusing temporal convolutions with spatial graph networks provides an edge in non-stationary financial markets."),
        ("fig_scatter_reg.png",
         "Fig. 8. Correlation analysis between predicted probabilities and actual forward returns.",
         "The scatter plot in Fig. 8 demonstrates the system's calibration. The regression curve indicates a positive correlation between the ensemble's predicted bullish probabilities and actual 5-day forward returns, suggesting that the output of the meta-learner aligns with future directional magnitude.")
    ],
    "C. Ablation Study": [
        ("fig_roc.png",
         "Fig. 9. Receiver Operating Characteristic (ROC) evaluating discriminative capability.",
         "The ROC curve analysis in Fig. 9 evaluates the discriminative capability of the full FinSight stacking ensemble. The Area Under the Curve (AUC) suggests that the integration of the Bayesian uncertainty veto and regime-adaptive loss functions contributes to the classification task performance."),
        ("fig_pr.png",
         "Fig. 10. Precision-Recall curves highlighting robustness in imbalanced regime states.",
         "The Precision-Recall curves in Fig. 10 emphasize the ensemble's performance under imbalanced regime states, demonstrating its ability to identify high-magnitude entry points while controlling false-positive rates."),
        ("fig_loss_curve.png",
         "Fig. 11. TCN-MHA Training vs. Validation Loss convergence across epochs.",
         "Fig. 11 illustrates the convergence properties of the TCN-MHA block. The synchronized decline of both training and validation loss indicates stable training dynamics.")
    ],
    "E. Regime-Adaptive Weight Distribution": [
        ("fig_regime.png",
         "Fig. 12. Dynamic equity curve mapped against detected market regimes.",
         "The time-series visualization in Fig. 12 maps the equity progression against dynamically detected macro-regimes, demonstrating the practical effect of the regime-adaptive asymmetric loss function in promoting defensive positions during Bear regimes.")
    ],
    "F. MC Uncertainty Quantification Results": [
        ("fig_confusion.png",
         "Fig. 13. Stacking Ensemble Confusion Matrix indicating asymmetric risk protection.",
         "Fig. 13 presents the ensemble's confusion matrix. The slight skew in off-diagonal misclassifications reflects the deliberate design of the regime-adaptive asymmetric loss, prioritizing capital preservation during bear markets."),
        ("fig_shap.png",
         "Fig. 14. SHAP Feature Importance analyzing multi-modal structural influence.",
         "Finally, the SHAP value analysis in Fig. 14 elucidates the LightGBM decision mechanics, indicating that GARCH volatility metrics and FinBERT sentiment proxies contribute meaningfully to the final predictive output alongside the foundational TCN embeddings.")
    ]
}

def main():
    print(f"Loading {DOC_PATH}...")
    doc = Document(DOC_PATH)
    
    
    insertions = []
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        for header, figs in MAPPINGS.items():
            if header in text:
                if i + 1 < len(doc.paragraphs):
                    insertions.append((doc.paragraphs[i+1], figs))
                else:
                    insertions.append((None, figs))
                break

    print(f"Found {len(insertions)} insertion points matching mapped sections.")

    for target_p, figs in insertions:
        for fname, caption, explanation in figs:
            img_path = os.path.join(PLOTS_DIR, fname)
            if not os.path.exists(img_path):
                print(f"  [!] Missing image: {fname}")
                continue
                
            if target_p is not None:
                p_img = target_p.insert_paragraph_before()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_path, width=Inches(3.0))
                
                p_cap = target_p.insert_paragraph_before()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = "Times New Roman"
                r_cap.font.size = Pt(9)
                r_cap.italic = True
                
                p_exp = target_p.insert_paragraph_before(explanation)
                p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_exp.runs[0].font.name = "Times New Roman"
                p_exp.runs[0].font.size = Pt(10)
                
                target_p.insert_paragraph_before("") # Spacing
            else:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_path, width=Inches(3.0))
                
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = "Times New Roman"
                r_cap.font.size = Pt(9)
                r_cap.italic = True
                
                p_exp = doc.add_paragraph(explanation)
                p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_exp.runs[0].font.name = "Times New Roman"
                p_exp.runs[0].font.size = Pt(10)
                
                doc.add_paragraph()
                
            print(f"  -> Inserted {fname} with enhanced novelty explanation.")

    doc.save(OUT_DOC)
    print(f"\nSuccessfully integrated refined novelty explanations and images into {OUT_DOC}")

if __name__ == "__main__":
    main()
